#!/usr/bin/env python3
"""KenDops 45-minute training deck."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUTS = [
    "/Users/swilliams/Documents/workspace/training/KENDOPS_DEVOPS_CLOUD_AI_AUTOMATION_45min_TRAINING.pptx",
    "/Users/swilliams/Library/Containers/net.whatsapp.WhatsApp/Data/tmp/documents/B7304750-91FD-4414-A567-00CE475E7717/KENDOPS_DEVOPS_CLOUD_AI_AUTOMATION_45min_TRAINING.pptx",
]

W, H = Inches(13.333), Inches(7.5)
BLACK = RGBColor(0x0B, 0x0C, 0x0E)
PAPER = RGBColor(0xF3, 0xF0, 0xE8)
INK = RGBColor(0x14, 0x16, 0x14)
BODY = RGBColor(0x3A, 0x3D, 0x38)
MUTED = RGBColor(0x6A, 0x6E, 0x66)
LIME = RGBColor(0xD6, 0xF2, 0x4A)
DEEP = RGBColor(0x1C, 0x3A, 0x28)
RULE = RGBColor(0xE2, 0xDD, 0xD2)
WHITE = RGBColor(0xF6, 0xF4, 0xEE)
SOFT = RGBColor(0xC4, 0xC8, 0xBE)

SERIF = "Iowan Old Style"
SANS = "Avenir Next"
MED = "Avenir Next Medium"
DEM = "Avenir Next Demi Bold"


def paint(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def style_run(run, text, font, size, color, bold=False, italic=False, spc=None):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = r_pr.makeelement(qn(tag), {})
            r_pr.append(node)
        node.set("typeface", font)
    if spc is not None:
        r_pr.set("spc", str(spc))


def add_text(slide, l, t, w, h, paragraphs, align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf._txBody.bodyPr.set("anchor", anchor)
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [para]
        if runs and runs[0].get("before"):
            p.space_before = Pt(runs[0]["before"])
        if runs and runs[0].get("after"):
            p.space_after = Pt(runs[0]["after"])
        if runs and runs[0].get("leading"):
            p.line_spacing = runs[0]["leading"]
        for spec in runs:
            run = p.add_run()
            style_run(
                run,
                spec.get("text", ""),
                spec.get("font", SANS),
                spec.get("size", 16),
                spec.get("color", INK),
                spec.get("bold", False),
                spec.get("italic", False),
                spec.get("spc"),
            )
    return box


def kicker(slide, l, t, w, text, color=DEEP, align=PP_ALIGN.LEFT):
    add_text(slide, l, t, w, Inches(0.28), [{
        "text": text.upper(), "font": MED, "size": 12, "color": color, "spc": 220,
    }], align=align)


def new(prs, color=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint(slide, color)
    return slide


def chrome(slide, page, dark=False):
    if dark:
        rect(slide, Inches(0), Inches(0), Inches(0.12), H, LIME)
        kicker(slide, Inches(0.62), Inches(0.36), Inches(4), "Kendops", LIME)
        add_text(slide, Inches(10.2), Inches(0.36), Inches(2.5), Inches(0.28), [{
            "text": f"{page:02d}", "font": SERIF, "size": 14, "color": SOFT,
        }], align=PP_ALIGN.RIGHT)
        return
    rect(slide, Inches(0), Inches(0), Inches(0.12), H, DEEP)
    kicker(slide, Inches(0.62), Inches(0.32), Inches(4), "Kendops")
    add_text(slide, Inches(10.2), Inches(0.32), Inches(2.5), Inches(0.28), [{
        "text": f"{page:02d}", "font": SERIF, "size": 14, "color": MUTED,
    }], align=PP_ALIGN.RIGHT)


def cover(prs):
    s = new(prs, BLACK)
    rect(s, Inches(0), Inches(0), Inches(0.18), H, LIME)
    kicker(s, Inches(0.72), Inches(0.48), Inches(6), "Kendops  ·  Training 2026", LIME)
    add_text(s, Inches(0.68), Inches(1.7), Inches(12), Inches(3.1), [
        {"text": "DevOps, cloud,", "font": SERIF, "size": 60, "color": WHITE, "leading": 0.92},
        {"text": "and AI automation.", "font": SERIF, "size": 60, "color": WHITE, "leading": 0.92},
    ])
    add_text(s, Inches(0.72), Inches(4.35), Inches(9), Inches(0.45), [{
        "text": "Forty-five minutes. One delivery path.",
        "font": SERIF, "size": 22, "color": LIME, "italic": True,
    }])
    add_text(s, Inches(0.72), Inches(6.35), Inches(8), Inches(0.35), [{
        "text": "BUILD     DEPLOY     OBSERVE     AUTOMATE",
        "font": MED, "size": 13, "color": SOFT, "spc": 120,
    }])
    add_text(s, Inches(7.4), Inches(6.35), Inches(5.3), Inches(0.35), [{
        "text": "CERTIFICATION AND CONSULTING",
        "font": MED, "size": 12, "color": SOFT, "spc": 80,
    }], align=PP_ALIGN.RIGHT)


def hour(prs):
    s = new(prs)
    chrome(s, 2)
    add_text(s, Inches(0.62), Inches(0.78), Inches(12), Inches(1.15), [
        {"text": "Leave with a workflow,", "font": SERIF, "size": 36, "color": INK, "leading": 0.98},
        {"text": "not a list of tools.", "font": SERIF, "size": 36, "color": INK, "leading": 0.98},
    ])
    outcomes = [
        ("01", "See the DevOps loop as shared ownership, not a job title."),
        ("02", "Place cloud and Terraform on that loop."),
        ("03", "Watch a change earn the right to ship: build, test, scan, deploy."),
        ("04", "Read the running system: metrics, logs, traces, and an alert with an owner."),
        ("05", "Use AI where it saves time, and keep a person on the risk."),
    ]
    for i, (n, text) in enumerate(outcomes):
        y = 2.02 + i * 0.48
        add_text(s, Inches(0.62), Inches(y), Inches(0.5), Inches(0.32), [{
            "text": n, "font": SERIF, "size": 15, "color": DEEP,
        }])
        add_text(s, Inches(1.25), Inches(y), Inches(11.2), Inches(0.38), [{
            "text": text, "font": SANS, "size": 17, "color": INK,
        }])
    moves = [
        ("01", "Foundations", "7 min"),
        ("02", "Cloud + IaC", "8 min"),
        ("03", "Pipeline", "8 min"),
        ("04", "Observe", "8 min"),
        ("05", "AI assist", "8 min"),
        ("06", "Live trace", "6 min"),
    ]
    rect(s, Inches(0.62), Inches(4.62), Inches(12.1), Pt(1.25), RULE)
    for i, (n, name, mins) in enumerate(moves):
        x = 0.62 + i * 2.05
        add_text(s, Inches(x), Inches(4.82), Inches(1.95), Inches(0.26), [{
            "text": n, "font": SERIF, "size": 14, "color": DEEP,
        }])
        add_text(s, Inches(x), Inches(5.12), Inches(1.95), Inches(0.32), [{
            "text": name, "font": DEM, "size": 14, "color": INK,
        }])
        add_text(s, Inches(x), Inches(5.48), Inches(1.95), Inches(0.28), [{
            "text": mins, "font": SANS, "size": 13, "color": MUTED,
        }])


def roles(prs):
    s = new(prs)
    chrome(s, 3)
    add_text(s, Inches(0.62), Inches(0.75), Inches(10), Inches(0.7), [{
        "text": "Four roles. One delivery.", "font": SERIF, "size": 36, "color": INK,
    }])
    items = [
        ("01", "DevOps engineer", "Builds the path from commit to release, and shortens the wait."),
        ("02", "Cloud engineer", "Designs infrastructure that can be locked down, scaled, and paid for."),
        ("03", "Platform engineer", "Turns that path into something the next team can use without a tour."),
        ("04", "AI automation engineer", "Removes the repeatable work, and leaves the judgment in place."),
    ]
    for i, (n, title, body) in enumerate(items):
        col = i % 2
        row = i // 2
        x = 0.62 + col * 6.3
        y = 1.85 + row * 2.25
        rect(s, Inches(x), Inches(y), Inches(0.42), Pt(2.5), DEEP)
        add_text(s, Inches(x), Inches(y + 0.15), Inches(1), Inches(0.3), [{
            "text": n, "font": SERIF, "size": 16, "color": DEEP,
        }])
        add_text(s, Inches(x), Inches(y + 0.5), Inches(5.6), Inches(0.45), [{
            "text": title, "font": SERIF, "size": 26, "color": INK,
        }])
        add_text(s, Inches(x), Inches(y + 1.1), Inches(5.6), Inches(0.8), [{
            "text": body, "font": SANS, "size": 16, "color": BODY,
        }])
    add_text(s, Inches(0.62), Inches(6.5), Inches(12), Inches(0.4), [{
        "text": "Growth shows up as projects other people can run, not as a longer list of tools.",
        "font": SERIF, "size": 16, "color": DEEP, "italic": True,
    }])


def system(prs):
    s = new(prs, BLACK)
    chrome(s, 4, dark=True)
    add_text(s, Inches(0.62), Inches(1.35), Inches(12), Inches(1.5), [
        {"text": "One workflow.", "font": SERIF, "size": 54, "color": WHITE, "leading": 0.95},
    ])
    words = ["Cloud", "DevOps", "Security", "AI", "Observe"]
    for i, word in enumerate(words):
        x = 0.62 + i * 2.45
        add_text(s, Inches(x), Inches(3.35), Inches(2.3), Inches(0.28), [{
            "text": f"0{i+1}", "font": SERIF, "size": 14, "color": LIME,
        }])
        add_text(s, Inches(x), Inches(3.7), Inches(2.3), Inches(0.5), [{
            "text": word, "font": SERIF, "size": 26, "color": WHITE,
        }])
    add_text(s, Inches(0.62), Inches(5.15), Inches(11), Inches(1.1), [{
        "text": "DevOps is how the team works. Cloud is where the system runs.\nSecurity decides what may ship. AI drafts the repeatable part.\nObservability tells you whether it stayed healthy.",
        "font": SANS, "size": 18, "color": SOFT, "leading": 1.2,
    }])


def loop(prs):
    s = new(prs)
    chrome(s, 5)
    add_text(s, Inches(0.62), Inches(0.75), Inches(12), Inches(0.9), [
        {"text": "Shorten the loop.", "font": SERIF, "size": 36, "color": INK, "leading": 1},
        {"text": "Do not skip the check.", "font": SERIF, "size": 36, "color": INK, "leading": 1},
    ])
    steps = [
        ("Plan", "Choose the smallest change."),
        ("Code", "Put it in Git."),
        ("Build", "Make it reproducible."),
        ("Test", "Prove the behavior."),
        ("Release", "Approve what passed."),
        ("Deploy", "Ship a known version."),
        ("Operate", "Run it on purpose."),
        ("Observe", "Metrics, logs, traces. Then plan again."),
    ]
    for i, (name, line) in enumerate(steps):
        col = i % 4
        row = i // 4
        x = 0.62 + col * 3.15
        y = 2.85 + row * 1.55
        add_text(s, Inches(x), Inches(y), Inches(2.9), Inches(0.28), [{
            "text": f"{i+1:02d}", "font": SERIF, "size": 14, "color": DEEP,
        }])
        add_text(s, Inches(x), Inches(y + 0.28), Inches(2.9), Inches(0.35), [{
            "text": name, "font": SERIF, "size": 22, "color": INK,
        }])
        add_text(s, Inches(x), Inches(y + 0.7), Inches(2.9), Inches(0.45), [{
            "text": line, "font": SANS, "size": 14, "color": BODY,
        }])


def tools(prs):
    s = new(prs)
    chrome(s, 6)
    add_text(s, Inches(0.62), Inches(0.72), Inches(10), Inches(0.5), [{
        "text": "The instruments", "font": SERIF, "size": 34, "color": INK,
    }])
    rows = [
        ("Cloud", "AWS, Azure, Google Cloud", "Compute, storage, network, identity."),
        ("Infrastructure", "Terraform, OpenTofu", "Reviewed changes, not clicked ones."),
        ("Delivery", "GitHub Actions, GitLab CI, Jenkins", "Build, test, release, deploy."),
        ("Containers", "Docker, Kubernetes", "The same application, anywhere it runs."),
        ("AI", "Model APIs, workflow tools", "Summaries, drafts, classification."),
        ("Observe", "Prometheus, Grafana, Alertmanager", "Metrics, logs, traces. A page with an owner."),
    ]
    rect(s, Inches(0.62), Inches(1.45), Inches(12.1), Pt(1.25), DEEP)
    for i, (name, tools_, job) in enumerate(rows):
        y = 1.6 + i * 0.82
        add_text(s, Inches(0.62), Inches(y), Inches(2.3), Inches(0.4), [{
            "text": name, "font": SERIF, "size": 20, "color": INK,
        }])
        add_text(s, Inches(3.1), Inches(y + 0.04), Inches(4.6), Inches(0.4), [{
            "text": tools_, "font": SANS, "size": 15, "color": BODY,
        }])
        add_text(s, Inches(7.8), Inches(y + 0.04), Inches(4.9), Inches(0.45), [{
            "text": job, "font": SANS, "size": 15, "color": INK,
        }])


def cloud(prs):
    s = new(prs)
    chrome(s, 7)
    add_text(s, Inches(0.62), Inches(0.72), Inches(12), Inches(0.55), [{
        "text": "Cloud, written down.", "font": SERIF, "size": 36, "color": INK,
    }])
    left = [
        "Capacity you can add, and remove, when demand moves.",
        "Identity, network rules, backups, and a budget.",
        "A managed service when running it yourself is not the work.",
    ]
    right = [
        "Infrastructure lives in Git, next to the application.",
        "A plan is reviewed before anyone is allowed to apply.",
        "State is protected. Secrets never enter the repository.",
    ]
    add_text(s, Inches(0.62), Inches(1.55), Inches(5.5), Inches(0.3), [{
        "text": "CLOUD", "font": MED, "size": 12, "color": DEEP, "spc": 180,
    }])
    add_text(s, Inches(6.9), Inches(1.55), Inches(5.5), Inches(0.3), [{
        "text": "INFRASTRUCTURE AS CODE", "font": MED, "size": 12, "color": DEEP, "spc": 160,
    }])
    for i, line in enumerate(left):
        y = 2.05 + i * 1.05
        add_text(s, Inches(0.62), Inches(y), Inches(0.4), Inches(0.3), [{
            "text": f"{i+1:02d}", "font": SERIF, "size": 14, "color": DEEP,
        }])
        add_text(s, Inches(1.1), Inches(y), Inches(5), Inches(0.85), [{
            "text": line, "font": SANS, "size": 16, "color": INK,
        }])
    for i, line in enumerate(right):
        y = 2.05 + i * 1.05
        add_text(s, Inches(6.9), Inches(y), Inches(0.4), Inches(0.3), [{
            "text": f"{i+1:02d}", "font": SERIF, "size": 14, "color": DEEP,
        }])
        add_text(s, Inches(7.4), Inches(y), Inches(5.2), Inches(0.85), [{
            "text": line, "font": SANS, "size": 16, "color": INK,
        }])
    rect(s, Inches(0.62), Inches(5.55), Inches(12.1), Pt(1.25), DEEP)
    add_text(s, Inches(0.62), Inches(5.75), Inches(12), Inches(0.7), [{
        "text": "Commit  →  plan  →  review  →  apply  →  verify",
        "font": SERIF, "size": 26, "color": INK,
    }])


def pipeline(prs):
    s = new(prs)
    chrome(s, 8)
    add_text(s, Inches(0.62), Inches(0.75), Inches(12), Inches(1.2), [
        {"text": "A change earns", "font": SERIF, "size": 40, "color": INK, "leading": 0.95},
        {"text": "the right to ship.", "font": SERIF, "size": 40, "color": INK, "leading": 0.95},
    ])
    steps = [
        ("01", "Build", "The same steps, every time."),
        ("02", "Test", "Behavior before opinion."),
        ("03", "Scan", "Secrets, dependencies, the image."),
        ("04", "Deploy", "A known version, to a named environment."),
        ("05", "Roll back", "The previous good release, on purpose."),
    ]
    for i, (n, title, body) in enumerate(steps):
        x = 0.62 + i * 2.5
        rect(s, Inches(x), Inches(3.15), Inches(0.55), Pt(3), LIME)
        add_text(s, Inches(x), Inches(3.4), Inches(2.3), Inches(0.3), [{
            "text": n, "font": SERIF, "size": 14, "color": DEEP,
        }])
        add_text(s, Inches(x), Inches(3.75), Inches(2.3), Inches(0.7), [{
            "text": title, "font": SERIF, "size": 24, "color": INK,
        }])
        add_text(s, Inches(x), Inches(4.55), Inches(2.3), Inches(0.9), [{
            "text": body, "font": SANS, "size": 15, "color": BODY,
        }])
    add_text(s, Inches(0.62), Inches(6.15), Inches(12), Inches(0.5), [{
        "text": "A red scan stops the merge. Waiving it is a decision, written down.",
        "font": SERIF, "size": 18, "color": DEEP, "italic": True,
    }])


def observe(prs):
    s = new(prs)
    chrome(s, 9)
    add_text(s, Inches(0.62), Inches(0.72), Inches(12), Inches(0.7), [{
        "text": "If you shipped it, you can see it.", "font": SERIF, "size": 34, "color": INK,
    }])
    pillars = [
        ("01", "Metrics", "Prometheus", "Rate, errors, and saturation. A number you can alert on."),
        ("02", "Logs", "ELK or Loki", "One request, in order, with enough context to act."),
        ("03", "Traces", "OpenTelemetry", "Where the time went, across every service it touched."),
    ]
    for i, (n, title, tool, body) in enumerate(pillars):
        x = 0.62 + i * 4.15
        add_text(s, Inches(x), Inches(1.7), Inches(3.8), Inches(0.28), [{
            "text": n, "font": SERIF, "size": 14, "color": DEEP,
        }])
        add_text(s, Inches(x), Inches(2.02), Inches(3.8), Inches(0.45), [{
            "text": title, "font": SERIF, "size": 28, "color": INK,
        }])
        add_text(s, Inches(x), Inches(2.55), Inches(3.8), Inches(0.3), [{
            "text": tool.upper(), "font": MED, "size": 12, "color": DEEP, "spc": 120,
        }])
        add_text(s, Inches(x), Inches(2.95), Inches(3.8), Inches(0.9), [{
            "text": body, "font": SANS, "size": 16, "color": BODY,
        }])
    rect(s, Inches(0.62), Inches(4.25), Inches(12.1), Pt(1.25), DEEP)
    practice = [
        ("Grafana", "The board the team looks at together."),
        ("Alertmanager", "Group, route, and silence. Page one owner."),
        ("Runbook", "The alert says what to do next."),
    ]
    for i, (title, body) in enumerate(practice):
        x = 0.62 + i * 4.15
        add_text(s, Inches(x), Inches(4.5), Inches(3.8), Inches(0.4), [{
            "text": title, "font": SERIF, "size": 20, "color": INK,
        }])
        add_text(s, Inches(x), Inches(5.0), Inches(3.8), Inches(0.7), [{
            "text": body, "font": SANS, "size": 15, "color": BODY,
        }])
    add_text(s, Inches(0.62), Inches(6.2), Inches(12), Inches(0.45), [{
        "text": "A green deploy is not a healthy service.",
        "font": SERIF, "size": 18, "color": DEEP, "italic": True,
    }])


def ai(prs):
    s = new(prs)
    chrome(s, 10)
    add_text(s, Inches(0.62), Inches(0.72), Inches(12), Inches(0.9), [
        {"text": "AI drafts. People decide.", "font": SERIF, "size": 36, "color": INK},
    ])
    uses = [
        ("Code", "Explain a failure. Draft a test. Write the note the next person needs."),
        ("Operations", "Group the logs. Name the likely cause. Do not page the room."),
        ("Service", "Sort the request. Draft the reply. A person sends it."),
        ("Release", "Write the notes from the commits that actually shipped."),
    ]
    rules = [
        "Secrets never go into a prompt.",
        "Production and destructive actions wait for approval.",
        "Generated code and commands are read before they run.",
        "Permissions are narrow. Actions are logged. Rollback exists.",
    ]
    for i, (title, body) in enumerate(uses):
        y = 1.75 + i * 0.85
        add_text(s, Inches(0.62), Inches(y), Inches(1.7), Inches(0.4), [{
            "text": title, "font": SERIF, "size": 20, "color": INK,
        }])
        add_text(s, Inches(2.4), Inches(y), Inches(5.3), Inches(0.7), [{
            "text": body, "font": SANS, "size": 15, "color": BODY,
        }])
    rect(s, Inches(8.15), Inches(1.75), Pt(1.5), Inches(3.5), DEEP)
    add_text(s, Inches(8.45), Inches(1.7), Inches(4.2), Inches(0.3), [{
        "text": "GUARDRAILS", "font": MED, "size": 12, "color": DEEP, "spc": 160,
    }])
    for i, rule in enumerate(rules):
        y = 2.15 + i * 0.8
        add_text(s, Inches(8.45), Inches(y), Inches(0.4), Inches(0.3), [{
            "text": f"{i+1:02d}", "font": SERIF, "size": 14, "color": DEEP,
        }])
        add_text(s, Inches(8.95), Inches(y), Inches(3.7), Inches(0.7), [{
            "text": rule, "font": SANS, "size": 14, "color": INK,
        }])


def demo(prs):
    s = new(prs, BLACK)
    chrome(s, 11, dark=True)
    add_text(s, Inches(0.62), Inches(0.8), Inches(12), Inches(0.7), [{
        "text": "Follow one change.", "font": SERIF, "size": 36, "color": WHITE,
    }])
    beats = [
        ("01", "Commit", "A developer pushes to Git."),
        ("02", "Build and test", "The pipeline installs, tests, and scans."),
        ("03", "Plan", "Infrastructure changes are previewed."),
        ("04", "Deploy", "An approved release reaches test."),
        ("05", "Observe", "Metrics, logs, and traces. The page names an owner."),
        ("06", "Assist", "AI drafts the summary. An engineer signs it."),
    ]
    for i, (n, title, body) in enumerate(beats):
        col = i % 3
        row = i // 3
        x = 0.62 + col * 4.15
        y = 1.85 + row * 1.7
        add_text(s, Inches(x), Inches(y), Inches(3.8), Inches(0.28), [{
            "text": n, "font": SERIF, "size": 14, "color": LIME,
        }])
        add_text(s, Inches(x), Inches(y + 0.3), Inches(3.8), Inches(0.4), [{
            "text": title, "font": SERIF, "size": 24, "color": WHITE,
        }])
        add_text(s, Inches(x), Inches(y + 0.8), Inches(3.8), Inches(0.55), [{
            "text": body, "font": SANS, "size": 15, "color": SOFT,
        }])
    add_text(s, Inches(0.62), Inches(5.7), Inches(12), Inches(0.9), [
        {"text": "Version the work. Automate the checks.", "font": SERIF, "size": 22, "color": LIME, "leading": 1.05},
        {"text": "Keep a human on the risk.", "font": SERIF, "size": 22, "color": LIME, "leading": 1.05},
    ])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "KenDops - DevOps, cloud, and AI automation"
    prs.core_properties.subject = "45-minute training"
    prs.core_properties.category = "Training"
    cover(prs)
    hour(prs)
    roles(prs)
    system(prs)
    loop(prs)
    tools(prs)
    cloud(prs)
    pipeline(prs)
    observe(prs)
    ai(prs)
    demo(prs)
    for path in OUTS:
        try:
            prs.save(path)
            print("saved", path)
        except OSError as err:
            print("skip", path, err)


if __name__ == "__main__":
    build()
