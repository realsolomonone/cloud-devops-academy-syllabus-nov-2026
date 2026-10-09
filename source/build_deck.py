#!/usr/bin/env python3
"""DevSecOps Academy syllabus deck. ASCII source only."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUTS = [
    "/Users/swilliams/Documents/workspace/training/DevSecOps_Academy_Syllabus.pptx",
    "/Users/swilliams/Documents/workspace/training/Cloud_DevOps_Academy_Engineering_Programme.pptx",
]

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x1C, 0x19, 0x17)
BODY = RGBColor(0x3F, 0x3A, 0x34)
MUTED = RGBColor(0x6B, 0x64, 0x5B)
CRIMSON = RGBColor(0xA5, 0x1C, 0x30)
PAPER = RGBColor(0xF6, 0xF3, 0xEE)
RULE = RGBColor(0xE3, 0xDB, 0xD0)
BLACK = RGBColor(0x14, 0x12, 0x11)
WHITE = RGBColor(0xF7, 0xF4, 0xEF)
SOFT = RGBColor(0xC8, 0xBE, 0xB4)

SERIF = "Iowan Old Style"
SANS = "Avenir Next"
MED = "Avenir Next Medium"
DEM = "Avenir Next Demi Bold"


def paint(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def style_run(run, text, font, size, color, bold=False, italic=False, spc=None):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = r_pr.makeelement(qn(tag), {})
            r_pr.append(node)
        node.set("typeface", font)
    if spc is not None:
        r_pr.set("spc", str(spc))


def add_text(slide, l, t, w, h, paragraphs, align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf._txBody.bodyPr.set("anchor", anchor)
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [para]
        if runs and runs[0].get("before"):
            p.space_before = Pt(runs[0]["before"])
        if runs and runs[0].get("after"):
            p.space_after = Pt(runs[0]["after"])
        if runs and runs[0].get("leading"):
            p.line_spacing = runs[0]["leading"]
        for spec in runs:
            run = p.add_run()
            style_run(
                run,
                spec.get("text", ""),
                spec.get("font", SANS),
                spec.get("size", 15),
                spec.get("color", INK),
                spec.get("bold", False),
                spec.get("italic", False),
                spec.get("spc"),
            )
    return box


def kicker(slide, l, t, w, text, color=CRIMSON, align=PP_ALIGN.LEFT):
    add_text(slide, l, t, w, Inches(0.28), [{
        "text": text.upper(), "font": MED, "size": 12, "color": color, "spc": 240,
    }], align=align)


def new(prs, color=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint(slide, color)
    return slide


def head(slide, right):
    rect(slide, Inches(0), Inches(0), Inches(0.12), H, CRIMSON)
    kicker(slide, Inches(0.55), Inches(0.32), Inches(6.5), "DevSecOps Academy")
    add_text(slide, Inches(7.2), Inches(0.32), Inches(5.55), Inches(0.28), [{
        "text": right.upper(), "font": MED, "size": 12, "color": MUTED, "spc": 180,
    }], align=PP_ALIGN.RIGHT)
    rect(slide, Inches(0.55), Inches(0.68), Inches(12.2), Pt(1), RULE)


def foot(slide, page, section):
    add_text(slide, Inches(0.55), Inches(7.05), Inches(1.2), Inches(0.28), [{
        "text": f"{page:02d}", "font": SERIF, "size": 13, "color": CRIMSON,
    }])
    add_text(slide, Inches(6.5), Inches(7.08), Inches(6.25), Inches(0.26), [{
        "text": section.upper(), "font": MED, "size": 11, "color": MUTED, "spc": 160,
    }], align=PP_ALIGN.RIGHT)


def cover(prs):
    s = new(prs, BLACK)
    rect(s, Inches(0), Inches(0), Inches(0.18), H, CRIMSON)
    kicker(s, Inches(0.75), Inches(0.48), Inches(8), "DevSecOps Academy", SOFT)
    add_text(s, Inches(0.72), Inches(1.85), Inches(12), Inches(2.6), [
        {"text": "Ship what you", "font": SERIF, "size": 60, "color": WHITE, "leading": 0.92},
        {"text": "can defend.", "font": SERIF, "size": 60, "color": WHITE, "leading": 0.92},
    ])
    add_text(s, Inches(0.75), Inches(4.35), Inches(9), Inches(0.7), [{
        "text": "A sixteen-week syllabus in cloud, delivery, and security.",
        "font": SERIF, "size": 22, "color": RGBColor(0xE7, 0xD3, 0xD6), "italic": True,
    }])
    rect(s, Inches(0.75), Inches(5.35), Inches(1.15), Pt(2.25), CRIMSON)
    add_text(s, Inches(0.75), Inches(5.6), Inches(11.5), Inches(0.85), [{
        "text": "Linux, Git, and GitLab CI.  AWS, Terraform, and Ansible.\nDocker, Kubernetes, and Argo CD.  Scans, Prometheus, Alertmanager, Grafana, and ELK.",
        "font": SANS, "size": 15, "color": SOFT, "leading": 1.25,
    }])


def charge(prs):
    s = new(prs)
    head(s, "The syllabus")
    add_text(s, Inches(0.55), Inches(0.95), Inches(12), Inches(1.5), [
        {"text": "Build nothing you cannot", "font": SERIF, "size": 36, "color": INK, "leading": 0.98},
        {"text": "scan, release, and explain.", "font": SERIF, "size": 36, "color": INK, "leading": 0.98},
    ])
    pillars = [
        ("01", "Practice", "Seventy percent of the hours are labs, projects, and review. Instruction earns its place by making the lab harder to fake."),
        ("02", "Judgement", "A working demo is not enough. Cost, failure, access, and the scan result are part of the grade."),
        ("03", "Proof", "The capstone is a system another engineer could operate from the runbook, without a private explanation."),
    ]
    for i, (n, title, body) in enumerate(pillars):
        x = 0.55 + i * 4.15
        rect(s, Inches(x), Inches(3.15), Inches(0.42), Pt(2.5), CRIMSON)
        add_text(s, Inches(x), Inches(3.4), Inches(3.8), Inches(0.3), [{
            "text": n, "font": SERIF, "size": 16, "color": CRIMSON,
        }])
        add_text(s, Inches(x), Inches(3.75), Inches(3.8), Inches(0.4), [{
            "text": title, "font": SERIF, "size": 26, "color": INK,
        }])
        add_text(s, Inches(x), Inches(4.35), Inches(3.85), Inches(1.6), [{
            "text": body, "font": SANS, "size": 15, "color": BODY,
        }])
    foot(s, 2, "The charge")


def outcomes(prs):
    s = new(prs)
    head(s, "Outcomes")
    add_text(s, Inches(0.55), Inches(0.88), Inches(10), Inches(0.55), [{
        "text": "By the end, a graduate can", "font": SERIF, "size": 32, "color": INK,
    }])
    items = [
        ("01", "Administer Linux and find a broken service from the logs, the port, and the firewall."),
        ("02", "Ship work through Git and GitLab, and stop a secret before it merges."),
        ("03", "Build a Docker image, scan it, and refuse to publish what the scan condemns."),
        ("04", "Provision AWS with least privilege and say what it costs."),
        ("05", "Express that infrastructure in Terraform and scan the plan before apply."),
        ("06", "Configure the host with Ansible so a second run changes nothing."),
        ("07", "Release to Kubernetes through GitLab CI and Argo CD, and roll a bad release back."),
        ("08", "Page on a real signal with Prometheus and Alertmanager, and investigate it in ELK."),
    ]
    for i, (n, text) in enumerate(items):
        col = i // 4
        row = i % 4
        x = 0.55 + col * 6.4
        y = 1.7 + row * 1.22
        add_text(s, Inches(x), Inches(y), Inches(0.55), Inches(0.35), [{
            "text": n, "font": SERIF, "size": 18, "color": CRIMSON,
        }])
        add_text(s, Inches(x + 0.6), Inches(y), Inches(5.5), Inches(0.9), [{
            "text": text, "font": SANS, "size": 16, "color": INK,
        }])
    foot(s, 3, "Outcomes")


def arc(prs):
    s = new(prs)
    head(s, "Sixteen weeks")
    add_text(s, Inches(0.55), Inches(0.86), Inches(8), Inches(0.5), [{
        "text": "The order of the work", "font": SERIF, "size": 32, "color": INK,
    }])
    months = [
        ("01", "Foundations", ["Linux and the shell", "Networking and SSH", "Git and GitLab", "AWS fundamentals"]),
        ("02", "Provision", ["Docker", "AWS architecture", "Terraform", "Ansible"]),
        ("03", "Delivery", ["GitLab CI", "Kubernetes", "Cluster operations", "The secure pipeline"]),
        ("04", "Operations", ["Metrics and alerts", "Logging with ELK", "Argo CD", "Capstone"]),
    ]
    for i, (n, name, weeks) in enumerate(months):
        x = 0.55 + i * 3.2
        add_text(s, Inches(x), Inches(1.7), Inches(2.9), Inches(0.28), [{
            "text": f"MONTH {n}", "font": MED, "size": 12, "color": CRIMSON, "spc": 180,
        }])
        add_text(s, Inches(x), Inches(2.02), Inches(2.9), Inches(0.45), [{
            "text": name, "font": SERIF, "size": 26, "color": INK,
        }])
        rect(s, Inches(x), Inches(2.6), Inches(1.0), Pt(2), CRIMSON)
        for j, title in enumerate(weeks):
            y = 2.9 + j * 0.85
            add_text(s, Inches(x), Inches(y), Inches(2.9), Inches(0.24), [{
                "text": f"{i * 4 + j + 1:02d}", "font": SERIF, "size": 13, "color": CRIMSON,
            }])
            add_text(s, Inches(x), Inches(y + 0.26), Inches(2.9), Inches(0.45), [{
                "text": title, "font": SANS, "size": 15, "color": INK,
            }])
    foot(s, 4, "Sixteen weeks")


def gates(prs):
    s = new(prs)
    head(s, "Security")
    add_text(s, Inches(0.55), Inches(0.86), Inches(12), Inches(0.5), [{
        "text": "Six gates. A red gate stops the merge.", "font": SERIF, "size": 30, "color": INK,
    }])
    items = [
        ("01", "Secrets", "Week 03", "A credential in the diff never reaches the default branch."),
        ("02", "Dependencies", "Week 09", "A known flaw in a library fails the pipeline."),
        ("03", "Static analysis", "Week 12", "SAST reviews the code on the merge, not after the incident."),
        ("04", "Images", "Weeks 05, 12", "The container is scanned before the registry will keep it."),
        ("05", "Infrastructure", "Week 07", "Terraform is scanned before anyone is allowed to apply."),
        ("06", "DAST and the record", "Week 12", "The running app is tested in a lower environment. The build records what shipped."),
    ]
    for i, (n, title, when, body) in enumerate(items):
        col = i % 3
        row = i // 3
        x = 0.55 + col * 4.2
        y = 1.75 + row * 2.45
        add_text(s, Inches(x), Inches(y), Inches(3.8), Inches(0.3), [{
            "text": n, "font": SERIF, "size": 16, "color": CRIMSON,
        }])
        add_text(s, Inches(x), Inches(y + 0.32), Inches(3.85), Inches(0.4), [{
            "text": title, "font": SERIF, "size": 22, "color": INK,
        }])
        add_text(s, Inches(x), Inches(y + 0.82), Inches(3.85), Inches(0.28), [{
            "text": when.upper(), "font": MED, "size": 12, "color": CRIMSON, "spc": 120,
        }])
        add_text(s, Inches(x), Inches(y + 1.16), Inches(3.85), Inches(0.9), [{
            "text": body, "font": SANS, "size": 15, "color": BODY,
        }])
    foot(s, 5, "Security gates")


def month(prs, page, label, title, weeks):
    s = new(prs)
    head(s, label)
    add_text(s, Inches(0.55), Inches(0.82), Inches(8), Inches(0.24), [{
        "text": label.upper(), "font": MED, "size": 12, "color": CRIMSON, "spc": 180,
    }])
    add_text(s, Inches(0.55), Inches(1.05), Inches(8), Inches(0.45), [{
        "text": title, "font": SERIF, "size": 32, "color": INK,
    }])
    positions = [(0.55, 1.75), (6.9, 1.75), (0.55, 4.25), (6.9, 4.25)]
    for (n, name, skills, evidence), (x, y) in zip(weeks, positions):
        rect(s, Inches(x), Inches(y), Inches(0.38), Pt(2.25), CRIMSON)
        add_text(s, Inches(x), Inches(y + 0.12), Inches(1.1), Inches(0.28), [{
            "text": n, "font": SERIF, "size": 14, "color": CRIMSON,
        }])
        add_text(s, Inches(x + 0.7), Inches(y + 0.08), Inches(5.0), Inches(0.36), [{
            "text": name, "font": SERIF, "size": 18, "color": INK,
        }])
        add_text(s, Inches(x), Inches(y + 0.55), Inches(5.85), Inches(0.75), [{
            "text": skills, "font": SANS, "size": 14, "color": BODY,
        }])
        add_text(s, Inches(x), Inches(y + 1.35), Inches(5.85), Inches(0.7), [[
            {"text": "EVIDENCE  ", "font": MED, "size": 11, "color": CRIMSON, "spc": 80},
            {"text": evidence, "font": SANS, "size": 14, "color": INK},
        ]])
    foot(s, page, title)


def capstone(prs):
    s = new(prs)
    head(s, "Week 16")
    add_text(s, Inches(0.55), Inches(0.86), Inches(10), Inches(0.5), [{
        "text": "The capstone", "font": SERIF, "size": 32, "color": INK,
    }])
    add_text(s, Inches(0.55), Inches(1.45), Inches(12), Inches(0.55), [{
        "text": "A production-style application, built by one person or a small team. The standard is a system a stranger could operate.",
        "font": SANS, "size": 16, "color": BODY,
    }])
    rows = [
        ("01", "Source", "GitLab. A branching model. Merge requests."),
        ("02", "Cloud", "Terraform on AWS. Least privilege, tags, a cost note."),
        ("03", "Pipeline", "GitLab CI tests, builds, scans, and publishes."),
        ("04", "Cluster", "Kubernetes, probes, limits, RBAC. EKS only if the budget allows."),
        ("05", "Release", "Argo CD from Git. Promotion and rollback, shown."),
        ("06", "Signals", "Prometheus, Grafana, Alertmanager. Logs in ELK."),
        ("07", "Gates", "A planted finding blocked the pipeline. The fix is in the history."),
        ("08", "Record", "Diagram, README, findings, runbook, and a live demonstration."),
    ]
    for i, (n, name, text) in enumerate(rows):
        col = 0 if i < 4 else 1
        row = i if i < 4 else i - 4
        x = 0.55 + col * 6.4
        y = 2.2 + row * 1.1
        add_text(s, Inches(x), Inches(y), Inches(0.5), Inches(0.3), [{
            "text": n, "font": SERIF, "size": 16, "color": CRIMSON,
        }])
        add_text(s, Inches(x + 0.55), Inches(y - 0.02), Inches(5.5), Inches(0.3), [{
            "text": name, "font": DEM, "size": 16, "color": INK,
        }])
        add_text(s, Inches(x + 0.55), Inches(y + 0.32), Inches(5.5), Inches(0.5), [{
            "text": text, "font": SANS, "size": 14, "color": BODY,
        }])
    foot(s, 10, "Capstone")


def assessment(prs):
    s = new(prs)
    head(s, "Assessment")
    add_text(s, Inches(0.55), Inches(0.86), Inches(10), Inches(0.5), [{
        "text": "What is graded", "font": SERIF, "size": 32, "color": INK,
    }])
    rows = [
        ("25%", "Weekly labs", "The work of the week, including the troubleshooting."),
        ("20%", "Projects and review", "Deliverables, writing, and the hygiene of the Git history."),
        ("15%", "Mid-programme", "Docker, AWS, Terraform, Ansible, and a pipeline that scans."),
        ("15%", "Security and signals", "Gates that fired, a fix, dashboards, logs, and alerts."),
        ("25%", "Capstone", "Architecture, automation, security, cost, and the demonstration."),
    ]
    y = 1.6
    for weight, name, detail in rows:
        add_text(s, Inches(0.55), Inches(y), Inches(1.5), Inches(0.45), [{
            "text": weight, "font": SERIF, "size": 26, "color": CRIMSON,
        }])
        add_text(s, Inches(2.2), Inches(y), Inches(3.4), Inches(0.4), [{
            "text": name, "font": SERIF, "size": 20, "color": INK,
        }])
        add_text(s, Inches(5.8), Inches(y + 0.04), Inches(6.8), Inches(0.4), [{
            "text": detail, "font": SANS, "size": 15, "color": BODY,
        }])
        y += 0.72
    add_text(s, Inches(0.55), Inches(5.4), Inches(12), Inches(1.2), [{
        "text": "Three sessions a week: the idea, the lab, then the review. Two to four hours between them. A claim without a repository or a demonstration does not count.",
        "font": SANS, "size": 15, "color": BODY,
    }])
    foot(s, 11, "Assessment")


def close(prs):
    s = new(prs, BLACK)
    rect(s, Inches(0), Inches(0), Inches(0.18), H, CRIMSON)
    kicker(s, Inches(0.75), Inches(2.15), Inches(8), "DevSecOps Academy", SOFT)
    add_text(s, Inches(0.72), Inches(2.6), Inches(12), Inches(2.2), [
        {"text": "Ship what you", "font": SERIF, "size": 54, "color": WHITE, "leading": 0.95},
        {"text": "can defend.", "font": SERIF, "size": 54, "color": WHITE, "leading": 0.95},
    ])
    add_text(s, Inches(0.75), Inches(5.15), Inches(10), Inches(0.4), [{
        "text": "Sixteen weeks.", "font": SERIF, "size": 20, "color": RGBColor(0xE7, 0xD3, 0xD6), "italic": True,
    }])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "DevSecOps Academy - Ship what you can defend"
    prs.core_properties.subject = "Sixteen-week syllabus"
    prs.core_properties.category = "Syllabus"
    cover(prs)
    charge(prs)
    outcomes(prs)
    arc(prs)
    gates(prs)
    month(prs, 6, "Weeks 01 - 04", "Foundations", [
        ("01", "Linux and the shell",
         "Filesystem, users, permissions, sudo, systemd, logs, and Bash.",
         "A least-privilege baseline and a scripted health report."),
        ("02", "Networking and SSH",
         "Addressing, DNS, TLS, firewalls, keys, patches, and a host baseline.",
         "Key-based access between two hosts, and a troubleshooting runbook."),
        ("03", "Git and GitLab",
         "Branches, merge requests, protected branches, and secret scanning.",
         "A reviewed project in which a planted secret is blocked."),
        ("04", "AWS fundamentals",
         "IAM, MFA, EC2, S3, VPC, security groups, budgets, and CloudTrail.",
         "A locked-down web server and an S3 site, with a cost and security note."),
    ])
    month(prs, 7, "Weeks 05 - 08", "Provision", [
        ("05", "Docker",
         "Images, multi-stage builds, Compose, volumes, and an image scan.",
         "An application and database. Critical findings fixed before anyone else runs it."),
        ("06", "AWS architecture",
         "Load balancing, Auto Scaling, Lambda, CloudWatch, private subnets, secrets.",
         "A highly available sample, and a Lambda with an alarm."),
        ("07", "Terraform",
         "Workflow, state, modules, and an infrastructure scan before apply.",
         "A modular network and compute stack, applied, then destroyed cleanly."),
        ("08", "Ansible",
         "Roles, Vault, idempotence, and a hardening pass after Terraform.",
         "Nginx from a role. The second run changes nothing."),
    ])
    month(prs, 8, "Weeks 09 - 12", "Delivery", [
        ("09", "GitLab CI",
         "Stages, runners, a registry, and a pipeline that fails closed on scans.",
         "A planted secret blocks the merge. A bad release is rolled back."),
        ("10", "Kubernetes",
         "Pods, Deployments, Services, probes, and a local cluster.",
         "A multi-replica app, then a broken rollout read from events and logs."),
        ("11", "Cluster operations",
         "Ingress, storage, Helm, network policy, RBAC, and security contexts.",
         "A Helm release running as a restricted service account."),
        ("12", "The secure pipeline",
         "SAST, dependencies, secrets, images, IaC, a DAST pass, and an SBOM.",
         "A planted finding stops the pipeline until the fix is in the history."),
    ])
    month(prs, 9, "Weeks 13 - 16", "Operations", [
        ("13", "Prometheus and alerts",
         "Scrape, PromQL, Grafana, and Alertmanager routing, grouping, and silences.",
         "Routes for CPU, disk, downtime, and repeated login failure."),
        ("14", "Logging with ELK",
         "Structured logs, Elasticsearch, Kibana, ingestion, and redaction.",
         "An investigation that names an error or an authentication failure."),
        ("15", "Argo CD",
         "Desired state, sync, drift, promotion, and rollback. Scans still apply.",
         "A release from Git, drift corrected, then promoted and rolled back."),
        ("16", "Capstone",
         "Architecture, security, cost, an incident, and the handover.",
         "The platform, the runbook, and a live demonstration."),
    ])
    capstone(prs)
    assessment(prs)
    close(prs)
    for path in OUTS:
        prs.save(path)
        print(path)


if __name__ == "__main__":
    build()
